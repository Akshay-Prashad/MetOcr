"""
Write extracted rows back into the target xlsx, matching the existing
column layout (verified against MO_9_2_029_c.xlsx). Writes into a copy
so the original is never mutated; flags low-confidence cells with a
comment-like marker so a human reviewer can find them fast.

Usage:
    python xlsx_writeback.py extracted_rows.json template.xlsx out.xlsx
"""

import json
import sys
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

from schema import LEFT_PAGE_SCHEMA, day_to_row

LOW_CONF_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")


def write_rows(rows: list[dict], template_path: str, out_path: str, sheet_name="Sheet1"):
    wb = openpyxl.load_workbook(template_path)
    ws = wb[sheet_name]

    flagged = []

    for row_data in rows:
        day = row_data.get("day")
        if day is None:
            print(f"WARNING: row missing 'day', skipping: {row_data}")
            continue

        xlsx_row = day_to_row(int(day))

        for field, (col_letter, desc, _type) in LEFT_PAGE_SCHEMA.items():
            if field not in row_data or row_data[field] is None:
                continue
            cell = ws[f"{col_letter}{xlsx_row}"]
            cell.value = row_data[field]

            conf_key = f"{field}_confidence"
            if row_data.get(conf_key) == "low":
                cell.fill = LOW_CONF_FILL
                cell.comment = Comment("Low-confidence VLM read -- verify against scan", "VLM pipeline")
                flagged.append(f"Day {day}, {desc} ({col_letter}{xlsx_row})")

    wb.save(out_path)
    print(f"Wrote {out_path}")
    if flagged:
        print(f"\n{len(flagged)} low-confidence cells flagged for review:")
        for f in flagged:
            print(f"  - {f}")


if __name__ == "__main__":
    json_path, template_path, out_path = sys.argv[1], sys.argv[2], sys.argv[3]
    with open(json_path) as f:
        rows = json.load(f)
    write_rows(rows, template_path, out_path)
